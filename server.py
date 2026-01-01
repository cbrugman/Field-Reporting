from flask import Flask, request, send_file, render_template, jsonify, send_from_directory
from flask_cors import CORS
from weasyprint import HTML, CSS
from staticmap import StaticMap, CircleMarker, IconMarker
import io
import json
import base64
from PIL import Image, ImageOps
import tempfile
import os
import re
import database
import matcher
from docx import Document
from docx.shared import Inches, Pt, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH

app = Flask(__name__)
CORS(app, resources={r"/*": {"origins": "*"}})  # Enable CORS for all routes with explicit wildcard

# Ensure DB is initialized
database.init_db()

def parse_species_list(s):
    if not s: 
        return {'entries': [], 'species_count': 0, 'individual_count': 0}
    
    # Split lines/commas
    raw_items = [x.strip() for x in s.replace('\r\n', '\n').replace(',', '\n').split('\n') if x.strip()]
    parsed_items = []
    total_individuals = 0
    
    for item_str in raw_items:
        # 1. Try "Name (5)"
        match = re.match(r"(.*)\s\((\d+)\)$", item_str)
        if match:
            name = match.group(1).strip()
            count = int(match.group(2))
        else:
            # 2. Try "Name 5" (no parens)
            match_no_parens = re.match(r"(.*)\s(\d+)$", item_str)
            if match_no_parens:
                name = match_no_parens.group(1).strip()
                count = int(match_no_parens.group(2))
            else:
                # 3. Try "5 Name" (Leading number)
                match_leading = re.match(r"^(\d+)\s+(.*)$", item_str)
                if match_leading:
                    count = int(match_leading.group(1))
                    name = match_leading.group(2).strip()
                else:
                    # 4. Default
                    name = item_str
                    count = 1
        
        parsed_items.append({'name': name, 'count': count})
        total_individuals += count

    return {
        'entries': parsed_items,
        'species_count': len(parsed_items),
        'individual_count': total_individuals
    }



@app.route('/')
def index():
    return send_file('Field Report Generator.html')

@app.route('/generator')
def generator():
    return send_file('Field Report Generator.html')

@app.route('/collector')
def collector():
    return send_file('field collector.html')

@app.route('/database')
def database_page():
    return send_file('CMNR species tracker.html')

@app.route('/api/verify-species', methods=['POST'])
def verify_species():
    try:
        print("Received verify-species request")
        """
        Receives JSON: { "data": { "plants": "...", "animals": "..." } }
        Returns analysis: 
        {
            "plants": [ { "original": "Blue Jya", "status": "fuzzy", "match": "Blue Jay" }, ... ]
        }
        """
        if not request.is_json:
            print("Error: Request is not JSON")
            return jsonify({"error": "Request must be JSON"}), 400

        req = request.json
        data = req.get('data', {})
        print(f"Processing data keys: {data.keys()}")
        
        # We only care about these 3 fields
        categories = ['plants', 'animals', 'fungi']
        
        response = {}
        
        # Map frontend fields to DB categories
        category_map = {
            'plants': ['Plants'],
            'fungi': ['Mushrooms'],
            'animals': ['Birds', 'Mammals', 'Insects', 'Herptiles']
        }
        
        for cat in categories:
            raw_text = data.get(cat, "")
            # Use our existing parser to get names
            parsed = parse_species_list(raw_text)
            
            # Determine allowed DB categories for this input field
            allowed_cats = category_map.get(cat)
            
            # Use dictionary to merge duplicates by Key (ID or Name)
            merged_results = {}

            for entry in parsed['entries']:
                name = entry['name']
                raw_count = entry['count']
                
                # Match logic
                try:
                    exact, fuzzy = matcher.find_matches(name, allowed_categories=allowed_cats)
                except Exception as match_err:
                    print(f"Matcher error for {name}: {match_err}")
                    raise match_err
                
                # Determine Key and Status
                if exact:
                    key = f"ID:{exact['id']}"
                    status = "exact"
                    match_name = exact['common_name']
                    match_id = exact['id']
                    score = 1.0
                elif fuzzy:
                    top = fuzzy[0]
                    key = f"ID:{top['id']}"
                    status = "fuzzy"
                    match_name = top['common_name']
                    match_id = top['id']
                    score = top['score']
                else:
                    # New species: Key by normalized name to merge "Sp A" and "Sp A"
                    norm_name = name.lower().strip()
                    key = f"NEW:{norm_name}"
                    status = "new"
                    match_name = None
                    match_id = None
                    score = 0

                # Merge Logic
                if key in merged_results:
                    existing = merged_results[key]
                    existing['count'] += raw_count
                    
                    # Upgrade logic: If we found a BETTER match for this ID (e.g. Exact vs previously Fuzzy), update details
                    # Or if we found a match for a previously "New" item (unlikely with this key logic, but good practice)
                    if status == 'exact' and existing['status'] != 'exact':
                         existing['status'] = 'exact'
                         existing['match'] = match_name
                         existing['score'] = 1.0
                         # Keep original name of the FIRST one? Or append? 
                         # Let's keep the original name of the BEST match or just the first encountered.
                         # For verification context, showing "Var A, Var B" might be noisy. Let's stick to first.
                else:
                    merged_results[key] = {
                        "original": name,
                        "status": status,
                        "match": match_name,
                        "id": match_id,
                        "count": raw_count,
                        "score": score
                    }

            # Convert map back to list
            response[cat] = list(merged_results.values())
            
        return jsonify(response)
    except Exception as e:
        import traceback
        traceback.print_exc()
        return jsonify({"error": str(e), "trace": traceback.format_exc()}), 500

@app.route('/api/add-visit', methods=['POST'])
def add_visit():
    try:
        if not request.is_json:
            return jsonify({"error": "Request must be JSON"}), 400
        
        req = request.json
        date = req.get('date')
        items = req.get('items', [])
        
        if not date:
            return jsonify({"error": "Missing date"}), 400
            
        added_count = 0
        
        for item in items:
            action = item.get('action')
            count = int(item.get('count', 1))
            category = item.get('category')
            original_name = item.get('original')
            
            # Categories mapping: frontend key -> DB value
            # Frontend uses 'plants', 'animals', 'fungi'
            # Backend should normalize. 'plants' -> 'Plants', etc.
            # But get_or_create_species allows passing category directly.
            # If the frontend provides a specific category, use it.
            # Otherwise, use the fallback map.
            cat_map = {
                'plants': 'Plants',
                'animals': 'Animals',
                'fungi': 'Fungi'
            }
            db_category = category if category else cat_map.get(item.get('category_key'), 'Uncategorized')
            scientific_name = item.get('scientific_name')
            
            species_id = None
            
            if action == 'verify':
                # Use existing ID
                species_id = item.get('match_id')
                if not species_id:
                     # Fallback if no ID provided (unlikely with UI logic)
                     species_id = database.get_or_create_species(item.get('match_name'), category=db_category)
                     
            elif action == 'create_new':
                # Create with original name
                species_id = database.get_or_create_species(original_name, scientific_name=scientific_name, category=db_category)
                
            if species_id:
                database.add_sighting(species_id, date, count, source='Web Report')
                added_count += 1
                
        return jsonify({"added": added_count})

    except Exception as e:
        import traceback
        traceback.print_exc()
        return jsonify({"error": str(e), "trace": traceback.format_exc()}), 500

@app.route('/api/species', methods=['GET'])
def get_species():
    try:
        species_list = database.get_all_species()
        return jsonify(species_list)
    except Exception as e:
        import traceback
        traceback.print_exc()
        return jsonify({"error": str(e)}), 500

@app.route('/api/visits', methods=['GET'])
def get_visits():
    try:
        visits = database.get_all_visits()
        return jsonify(visits)
    except Exception as e:
        import traceback
        traceback.print_exc()
        return jsonify({"error": str(e)}), 500

@app.route('/api/visits/<date>', methods=['GET'])
def get_visit_details_route(date):
    try:
        details = database.get_visit_details(date)
        return jsonify(details)
    except Exception as e:
        import traceback
        traceback.print_exc()
        return jsonify({"error": str(e)}), 500

@app.route('/api/visits/<date>', methods=['DELETE'])
def delete_visit_route(date):
    try:
        count = database.delete_visit(date)
        return jsonify({"deleted": count})
    except Exception as e:
        import traceback
        traceback.print_exc()
        return jsonify({"error": str(e)}), 500

@app.route('/api/deduplicate', methods=['POST'])
def run_deduplication():
    try:
        count = database.deduplicate_sightings()
        return jsonify({"message": f"Cleaned up {count} duplicate records."})
    except Exception as e:
        import traceback
        traceback.print_exc()
        return jsonify({"error": str(e)}), 500

@app.route('/api/normalize-dates', methods=['POST'])
def normalize_dates_route():
    try:
        updated = database.normalize_dates()
        # Also re-run deduplication since merging dates might create duplicates
        deduped = database.deduplicate_sightings()
        return jsonify({
            "message": f"Normalized {updated} dates. Cleaned up {deduped} resulting duplicates."
        })
    except Exception as e:
        import traceback
        traceback.print_exc()
        return jsonify({"error": str(e)}), 500

def generate_map_from_photos(photos):
    map_image_url = None
    temp_files = [] # Keep track of temp files to delete
    try:
        with open("error_log.txt", "a") as log:
            log.write("--- Map Generation Start ---\n")
            
        # Create a map with size 800x600
        m = StaticMap(800, 600)
        has_markers = False
        
        for p in photos:
            gps = p.get('gps', '')
            # with open("error_log.txt", "a") as log:
            #     log.write(f"Checking Photo GPS: '{gps}'\n")

            if gps and ',' in gps:
                try:
                    parts = gps.split(',')
                    lat = float(parts[0].strip())
                    lon = float(parts[1].strip())
                    
                    # Process image for marker
                    if 'dataUrl' in p:
                        img_data = p['dataUrl'].split(',')[1]
                        img_bytes = base64.b64decode(img_data)
                        
                        with Image.open(io.BytesIO(img_bytes)) as pil_img:
                            # Convert to RGB
                            pil_img = pil_img.convert("RGB")
                            pil_img.thumbnail((64, 64))
                            
                            # Add "Photo Frame" effect
                            pil_img = ImageOps.expand(pil_img, border=3, fill='white')
                            pil_img = ImageOps.expand(pil_img, border=1, fill='#555555')
                            pil_img = pil_img.convert("RGBA")
                            
                            # Save to temp file
                            tf = tempfile.NamedTemporaryFile(suffix='.png', delete=False)
                            pil_img.save(tf, format='PNG')
                            tf.close()
                            temp_files.append(tf.name)
                            
                            # Center marker
                            w, h = pil_img.size
                            offset_x = -w // 2
                            offset_y = -h // 2
                            
                            marker = IconMarker((lon, lat), tf.name, offset_x, offset_y)
                            m.add_marker(marker)
                            has_markers = True
                            
                    else:
                         marker = CircleMarker((lon, lat), 'red', 10)
                         m.add_marker(marker)
                         has_markers = True
                         
                except Exception as ex:
                    with open("error_log.txt", "a") as log:
                        log.write(f"Failed to process marker '{gps}': {ex}\n")
                    continue 
        
        if has_markers:
            try:
                image = m.render()
                img_buffer = io.BytesIO()
                image.save(img_buffer, format='PNG')
                img_str = base64.b64encode(img_buffer.getvalue()).decode('utf-8')
                map_image_url = f"data:image/png;base64,{img_str}"
            except Exception as render_ex:
                 with open("error_log.txt", "a") as log:
                    log.write(f"Render failed: {render_ex}\n")
        else:
             with open("error_log.txt", "a") as log:
                log.write("No valid markers found.\n")

    except Exception as e:
        with open("error_log.txt", "a") as log:
            log.write(f"Map generation error (outer): {e}\n")
        print(f"Map generation error (outer): {e}")
        
    finally:
        # Clean up temp files
        for tf_path in temp_files:
            try:
                os.remove(tf_path)
            except:
                pass
                
    return map_image_url

@app.route('/generate-pdf', methods=['POST'])
def generate_pdf():
    try:
        # Get the JSON data from the request
        if not request.is_json:
            return "Expected JSON data", 400
            
        data = request.json
        
        # Pre-process lists (split strings into parsed objects for Jinja)
        # The frontend sends "Name (Count)" or just "Name".


        data['plants_data'] = parse_species_list(data.get('plants', ''))
        data['animals_data'] = parse_species_list(data.get('animals', ''))
        data['fungi_data'] = parse_species_list(data.get('fungi', ''))
        
        if 'photos' not in data or not isinstance(data['photos'], list):
            data['photos'] = []
        
        # Convert photoCols to int
        try:
            data['photoCols'] = int(data.get('photoCols', 3))
        except:
            data['photoCols'] = 3

        # Generate Map if GPS data exists
        map_image_url = generate_map_from_photos(data['photos'])
            
        data['map_image'] = map_image_url

        # Render HTML from template
        html_content = render_template('report_template.html', **data)
        
        # Debug: Save rendered HTML to inspect if needed using a flag or just print length
        # with open("debug_render.html", "w", encoding="utf-8") as f: f.write(html_content)

        # Create a file-like object
        pdf_io = io.BytesIO()

        # Generate PDF using WeasyPrint
        HTML(string=html_content, base_url=".").write_pdf(target=pdf_io)

        pdf_io.seek(0)

        # Send the file
        return send_file(
            pdf_io,
            mimetype='application/pdf',
            as_attachment=True,
            download_name='report.pdf' 
        )

    except Exception as e:
        import traceback
        error_msg = traceback.format_exc()
        print(f"Error generating PDF: {error_msg}")
        with open("error_log.txt", "w") as f:
            f.write(error_msg)
        return f"Server Error: {str(e)}", 500

@app.route('/generate-word', methods=['POST'])
def generate_word():
    try:
        if not request.is_json:
            return "Expected JSON data", 400
            
        data = request.json
        
        # 1. Parse Data
        data['plants_data'] = parse_species_list(data.get('plants', ''))
        data['animals_data'] = parse_species_list(data.get('animals', ''))
        data['fungi_data'] = parse_species_list(data.get('fungi', ''))
        
        if 'photos' not in data or not isinstance(data['photos'], list):
            data['photos'] = []

        # 2. Generate Map
        map_image_url = generate_map_from_photos(data['photos'])

        # 3. Create Word Document
        doc = Document()
        
        # Styles
        style = doc.styles['Normal']
        font = style.font
        font.name = 'Calibri'
        font.size = Pt(11)
        
        # Title
        title = doc.add_heading('Nature Reserve Visit Report', 0)
        title.alignment = WD_ALIGN_PARAGRAPH.CENTER
        
        
        # Metadata Section
        table = doc.add_table(rows=1, cols=2)
        table.autofit = True
        
        row_cells = table.add_row().cells
        p = row_cells[0].paragraphs[0]
        p.add_run("Reserve: ").bold = True
        p.add_run(f"{data.get('reserve', '')}")
        p.paragraph_format.space_after = Pt(0)
        
        p = row_cells[1].paragraphs[0]
        p.add_run("Date: ").bold = True
        p.add_run(f"{data.get('date', '')}")
        p.paragraph_format.space_after = Pt(0)
        
        row_cells = table.add_row().cells
        p = row_cells[0].paragraphs[0]
        p.add_run("Observers: ").bold = True
        p.add_run(f"{data.get('observers', '')}")
        p.paragraph_format.space_after = Pt(0)
        row_cells[1].text = ""

        # Weather Section
        doc.add_heading('Weather', level=2)
        w_table = doc.add_table(rows=1, cols=3)
        w_cells = w_table.rows[0].cells
        
        p = w_cells[0].paragraphs[0]
        p.add_run("Temp: ").bold = True
        p.add_run(f"{data.get('temp', '')} °C")
        p.paragraph_format.space_after = Pt(0)
        
        p = w_cells[1].paragraphs[0]
        p.add_run("Sky: ").bold = True
        p.add_run(f"{data.get('sky', '')}")
        p.paragraph_format.space_after = Pt(0)
        
        p = w_cells[2].paragraphs[0]
        p.add_run("Precip: ").bold = True
        p.add_run(f"{data.get('precip', '')}")
        p.paragraph_format.space_after = Pt(0)
        
        # Time Section
        t_table = doc.add_table(rows=1, cols=2)
        t_cells = t_table.rows[0].cells
        
        p = t_cells[0].paragraphs[0]
        p.add_run("Start Time: ").bold = True
        p.add_run(f"{data.get('start', '')}")
        p.paragraph_format.space_after = Pt(0)
        
        p = t_cells[1].paragraphs[0]
        p.add_run("End Time: ").bold = True
        p.add_run(f"{data.get('end', '')}")
        p.paragraph_format.space_after = Pt(0)

        # Text Sections
        for section, title in [('activities', 'Activities'), ('trails', 'State of Trails'), ('anthropogenic', 'Anthropogenic')]:
            doc.add_heading(title, level=2)
            doc.add_paragraph(data.get(section, '') or "None")

        # Species Sections
        for key, title in [('plants_data', 'Plants'), ('animals_data', 'Animals'), ('fungi_data', 'Fungi')]:
            doc.add_heading(title, level=2)
            entries = data[key]['entries']
            if not entries:
                doc.add_paragraph("None observed.")
            else:
                # Use a table for 3-column layout
                table = doc.add_table(rows=0, cols=3)
                table.autofit = True
                
                for i in range(0, len(entries), 3):
                    row_cells = table.add_row().cells
                    batch = entries[i:i+3]
                    for idx, entry in enumerate(batch):
                        row_cells[idx].text = f"{entry['name']} ({entry['count']})"
                        row_cells[idx].paragraphs[0].paragraph_format.space_after = Pt(0)

        # Photos Section
        doc.add_heading('Photos', level=2)
        
        # Handle User Photos
        if data['photos']:
            doc.add_heading('Field Photos', level=3)
            
            # Use a table for layout (2 columns)
            table = doc.add_table(rows=0, cols=2)
            table.autofit = True
            
            # Filter valid photos
            valid_photos = [p for p in data['photos'] if 'dataUrl' in p]
            
            # Iterate in pairs
            for i in range(0, len(valid_photos), 2):
                row_cells = table.add_row().cells
                batch = valid_photos[i:i+2]
                
                for idx, p in enumerate(batch):
                    cell = row_cells[idx]
                    try:
                        img_data = p['dataUrl'].split(',')[1]
                        img_bytes = base64.b64decode(img_data)
                        img_stream = io.BytesIO(img_bytes)
                        
                        # Add image to cell
                        paragraph = cell.paragraphs[0]
                        run = paragraph.add_run()
                        run.add_picture(img_stream, width=Cm(7.5))
                        
                        # Add caption
                        caption = []
                        if p.get('title'): caption.append(p['title'])
                        if p.get('notes'): caption.append(p['notes'])
                        
                        if caption:
                            paragraph.add_run("\n" + "\n".join(caption))
                            
                    except Exception as e:
                        print(f"Error adding photo to docx: {e}")
                        cell.text = "[Error]"

        # Handle Map
        if map_image_url:
            doc.add_page_break()
            doc.add_heading('Location Map', level=3)
            # Decode base64 map
            try:
                map_data = map_image_url.split(',')[1]
                map_bytes = base64.b64decode(map_data)
                map_stream = io.BytesIO(map_bytes)
                doc.add_picture(map_stream, width=Inches(6))
                doc.add_paragraph("GPS locations of photos.")
            except Exception as e:
                doc.add_paragraph(f"[Error adding map: {e}]")


        # Save to buffer
        docx_io = io.BytesIO()
        doc.save(docx_io)
        docx_io.seek(0)
        
        return send_file(
            docx_io,
            mimetype='application/vnd.openxmlformats-officedocument.wordprocessingml.document',
            as_attachment=True,
            download_name='report.docx'
        )

    except Exception as e:
        import traceback
        traceback.print_exc()
        return jsonify({"error": str(e), "trace": traceback.format_exc()}), 500

if __name__ == '__main__':
    port = int(os.environ.get('PORT', 5000))
    app.run(host='0.0.0.0', port=port)
